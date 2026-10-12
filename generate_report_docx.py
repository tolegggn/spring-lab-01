import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Page Margins: 1 inch (2.54 cm) all around
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
COLOR_PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)      # Deep Navy
COLOR_SECONDARY = RGBColor(0x25, 0x63, 0xEB)    # Royal Blue
COLOR_TEXT_MAIN = RGBColor(0x1F, 0x29, 0x37)    # Charcoal
COLOR_MUTED = RGBColor(0x4B, 0x55, 0x63)        # Gray

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_paragraph(text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_PRIMARY
    elif level == 2:
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_body_paragraph(text, bold_prefix=None, space_after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = COLOR_TEXT_MAIN
        r_pre.bold = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT_MAIN
    return p

def add_code_block(code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="3B82F6"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.0)
    run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    # Spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_screenshot_figure(image_path, caption_text):
    if os.path.exists(image_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run = p_img.add_run()
        run.add_picture(image_path, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

# ==================== COVER / METADATA HEADER ====================
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_inst.paragraph_format.space_after = Pt(2)
r = p_inst.add_run("JSC «INTERNATIONAL INFORMATION TECHNOLOGY UNIVERSITY»")
r.font.name = 'Calibri'
r.font.size = Pt(13)
r.bold = True
r.font.color.rgb = COLOR_PRIMARY

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(12)
r = p_sub.add_run("Faculty of Business, Media and Management · Department of Information Systems")
r.font.name = 'Calibri'
r.font.size = Pt(10.5)
r.font.color.rgb = COLOR_MUTED

p_lab = doc.add_paragraph()
p_lab.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_lab.paragraph_format.space_after = Pt(4)
r = p_lab.add_run("LABORATORY REPORT No. 6")
r.font.name = 'Calibri'
r.font.size = Pt(18)
r.bold = True
r.font.color.rgb = COLOR_PRIMARY

p_topic = doc.add_paragraph()
p_topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_topic.paragraph_format.space_after = Pt(14)
r = p_topic.add_run("Forms and Validation: Thymeleaf Templates, Form Object Binding, Bean Validation, Error Messages, Centralised Exception Handling")
r.font.name = 'Calibri'
r.font.size = Pt(11.5)
r.bold = True
r.font.color.rgb = COLOR_SECONDARY

# Metadata Table
meta_table = doc.add_table(rows=6, cols=2)
meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Course:", "RWPSF 3305 — Web Application Development with Spring Framework"),
    ("Educational Programme:", "6B06105 — Information Systems, year 3, semester 6"),
    ("Student:", "Tolegen Alikhan"),
    ("Group:", "IT1-2406 (2406IT)"),
    ("Variant:", "Variant 2 (Class-level constraint: reprintYear >= year, global errors)"),
    ("Lecturer:", "Yemberdiyeva Aknur, Senior Lector")
]

for i, (k, v) in enumerate(meta_data):
    row = meta_table.rows[i]
    cell_k, cell_v = row.cells[0], row.cells[1]
    cell_k.width = Inches(2.2)
    cell_v.width = Inches(4.3)
    set_cell_background(cell_k, "F1F5F9")
    set_cell_background(cell_v, "FFFFFF")
    set_cell_margins(cell_k, top=70, bottom=70, left=120, right=120)
    set_cell_margins(cell_v, top=70, bottom=70, left=120, right=120)
    
    p0 = cell_k.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    r0 = p0.add_run(k)
    r0.font.name = 'Calibri'
    r0.font.size = Pt(9.5)
    r0.bold = True
    r0.font.color.rgb = COLOR_TEXT_MAIN
    
    p1 = cell_v.paragraphs[0]
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(v)
    r1.font.name = 'Calibri'
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = COLOR_TEXT_MAIN

# Divider
p_div = doc.add_paragraph()
p_div.paragraph_format.space_before = Pt(10)
p_div.paragraph_format.space_after = Pt(10)
r = p_div.add_run("―" * 60)
r.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

# ==================== 1. AIM OF THE WORK ====================
add_header_paragraph("1. Aim of the Work and Repository Link", level=1)
add_body_paragraph(
    "The primary objective of this laboratory work is to construct a production-ready, interactive web form utilizing "
    "Spring MVC, Thymeleaf, and Jakarta Bean Validation. The key architectural achievements include:\n"
    "1. Form Object Binding: Binding an HTTP POST payload to a dedicated BookForm backing object, maintaining typed attributes independently from immutable domain entities.\n"
    "2. Declarative Bean Validation: Utilizing standard Jakarta constraints (@NotBlank, @Size, @Min, @Max, @NotNull) and handling validation results non-destructively through BindingResult.\n"
    "3. User Experience & PRG Pattern: Redisplaying validation errors alongside the corresponding inputs with retained user input upon failure, and enforcing the Post/Redirect/Get (PRG) pattern with flash attributes upon successful submission.\n"
    "4. Internationalization (i18n): Centralizing UI labels, form captions, and validation error messages into localized resource bundles (English messages.properties and Russian messages_ru.properties) with dynamic locale switching via ?lang=..\n"
    "5. Custom Constraint Development: Authoring a custom field-level validator (@NoDigits) ensuring author names contain no numerical digits.\n"
    "6. Centralised Exception Handling: Directing application-wide exceptions using targeted @ControllerAdvice (producing customized HTML error pages) and @RestControllerAdvice (producing RFC 7807/9457 ProblemDetail JSON payloads).\n"
    "7. Individual Assignment (Variant 2): Designing and implementing a class-level validation constraint (@ValidReprintYear) ensuring reprintYear >= year, surfaced via Thymeleaf's #fields.globalErrors()."
)

add_body_paragraph("Repository Branch: ", bold_prefix="Git Link: ")
add_body_paragraph("git checkout lab06 | Commit: lab06: form binding, validation, i18n, controller advice, variant 2")

# ==================== 2. CODE LISTINGS ====================
add_header_paragraph("2. Implementation and Core Code Listings", level=1)

add_header_paragraph("2.1. The Form-Backing Object (BookForm.java)", level=2)
add_body_paragraph(
    "The BookForm class acts as the Data Transfer Object (DTO) for form binding. It decouples browser interaction from the internal Book record, "
    "allowing partial binding and preserving invalid user inputs during re-rendering."
)
add_code_block("""package kz.iitu.springlab.web;

import jakarta.validation.constraints.*;
import kz.iitu.springlab.catalog.Genre;
import kz.iitu.springlab.web.validation.NoDigits;
import kz.iitu.springlab.web.validation.ValidReprintYear;

@ValidReprintYear
public class BookForm {

    @NotBlank(message = "{book.title.required}")
    @Size(max = 120, message = "{book.title.size}")
    private String title;

    @NotBlank(message = "{book.author.required}")
    @NoDigits
    private String author;

    @Min(value = 1450, message = "{book.year.range}")
    @Max(value = 2100, message = "{book.year.range}")
    private int year;

    // Individual assignment: Variant 2
    @Min(value = 1450, message = "{book.reprintYear.range}")
    @Max(value = 2100, message = "{book.reprintYear.range}")
    private Integer reprintYear;

    @NotNull(message = "{book.genre.required}")
    private Genre genre;

    private boolean available;

    public BookForm() {}

    // Getters and Setters for all fields...
}""")

add_header_paragraph("2.2. Custom Constraints: Field-Level and Class-Level", level=2)
add_body_paragraph(
    "Field constraint @NoDigits validates that an author name contains no digits. "
    "Class-level constraint @ValidReprintYear (Variant 2) validates cross-field logic: reprintYear >= year."
)
add_code_block("""// @NoDigits Validator
public class NoDigitsValidator implements ConstraintValidator<NoDigits, String> {
    @Override
    public boolean isValid(String value, ConstraintValidatorContext ctx) {
        if (value == null) return true; // Null/empty check is delegated to @NotBlank
        return value.chars().noneMatch(Character::isDigit);
    }
}

// Variant 2: @ValidReprintYear Validator (Class-Level)
public class ValidReprintYearValidator implements ConstraintValidator<ValidReprintYear, BookForm> {
    @Override
    public boolean isValid(BookForm form, ConstraintValidatorContext context) {
        if (form == null || form.getReprintYear() == null) return true;
        return form.getReprintYear() >= form.getYear();
    }
}""")

add_header_paragraph("2.3. The Form Controller Handlers (BookPageController.java)", level=2)
add_body_paragraph(
    "The controller wires the GET form display and POST acceptance. BindingResult is placed immediately following the @Valid BookForm."
)
add_code_block("""@GetMapping("/new")
public String newForm(Model model) {
    model.addAttribute("form", new BookForm());
    model.addAttribute("genres", Genre.values());
    return "books/form";
}

@PostMapping
public String create(@Valid @ModelAttribute("form") BookForm form,
                     BindingResult binding,          // IMMEDIATELY AFTER @Valid!
                     Model model,
                     RedirectAttributes redirect) {

    if (binding.hasErrors()) {
        model.addAttribute("genres", Genre.values()); // Re-populate reference data
        return "books/form";                          // Redisplay with errors
    }

    Book saved = service.create(new Book(null, form.getTitle(), form.getAuthor(), form.getYear()));
    redirect.addFlashAttribute("message", "Book «" + saved.title() + "» has been added");
    return "redirect:/books";                         // Post/Redirect/Get pattern
}""")

add_header_paragraph("2.4. Centralised Exception Handling (@ControllerAdvice)", level=2)
add_body_paragraph(
    "Two scoped advices differentiate browser error rendering (HTML error template) from REST client responses (RFC 7807 ProblemDetail)."
)
add_code_block("""@ControllerAdvice(annotations = Controller.class)
public class PageExceptionHandler {
    @ExceptionHandler(BookNotFoundException.class)
    public String handleNotFound(BookNotFoundException ex, Model model) {
        model.addAttribute("message", ex.getMessage());
        return "error/not-found";
    }
}

@RestControllerAdvice(annotations = RestController.class)
public class ApiExceptionHandler {
    @ExceptionHandler(BookNotFoundException.class)
    @ResponseStatus(HttpStatus.NOT_FOUND)
    public ProblemDetail handleNotFound(BookNotFoundException ex) {
        ProblemDetail problem = ProblemDetail.forStatusAndDetail(HttpStatus.NOT_FOUND, ex.getMessage());
        problem.setTitle("Book not found");
        return problem;
    }
}""")

# ==================== 3. RESOURCE BUNDLES ====================
add_header_paragraph("3. Resource Bundles (Internationalization)", level=1)
add_body_paragraph("Default bundle (messages.properties) and Russian bundle (messages_ru.properties):")
add_code_block("""# src/main/resources/messages.properties
book.new=New book
book.title=Title
book.author=Author
book.year=Year
book.reprintYear=Reprint Year
book.genre=Genre
form.submit=Save
book.title.required=The title must not be empty
book.title.size=The title must not exceed 120 characters
book.author.required=The author must not be empty
book.author.nodigits=Author name must not contain digits
book.year.range=The year must be between 1450 and 2100
book.reprintYear.range=The reprint year must be between 1450 and 2100
book.reprintYear.invalid=A reprint year must not precede the first edition year
book.genre.required=Choose a genre
book.available=Available
book.back=Back to catalogue
error.notfound.title=Book Not Found

# src/main/resources/messages_ru.properties
book.new=Новая книга
book.title=Название
book.author=Автор
book.year=Год
book.reprintYear=Год переиздания
book.genre=Жанр
form.submit=Сохранить
book.title.required=Название не может быть пустым
book.title.size=Название не длиннее 120 символов
book.author.required=Автор не может быть пустым
book.author.nodigits=Имя автора не должно содержать цифры
book.year.range=Год должен быть в диапазоне 1450–2100
book.reprintYear.range=Год переиздания должен быть в диапазоне 1450–2100
book.reprintYear.invalid=Год переиздания не может быть раньше года первого издания
book.genre.required=Выберите жанр
book.available=В наличии
book.back=Назад к каталогу
error.notfound.title=Книга не найдена""")

# ==================== 4. SCREENSHOTS AND DEMONSTRATION ====================
add_header_paragraph("4. Visual Verification and Demonstration Screenshots", level=1)

add_screenshot_figure("screenshots/01_form_initial_en.png", "Figure 1: Initial empty form in English (GET /books/new?lang=en)")
add_screenshot_figure("screenshots/02_form_initial_ru.png", "Figure 2: Initial empty form in Russian (GET /books/new?lang=ru)")
add_screenshot_figure("screenshots/03_form_empty_ru.png", "Figure 3: Submission of empty form in Russian, displaying localized error summary and per-field messages")
add_screenshot_figure("screenshots/04_year_range_no_genre.png", "Figure 4: Invalid year (1200) and unselected genre rejected while typed inputs survive")
add_screenshot_figure("screenshots/05_author_nodigits_violation.png", "Figure 5: Rejection of author containing numerical digits («Bloch 2018») via custom @NoDigits validator")
add_screenshot_figure("screenshots/06_variant2_class_level_error.png", "Figure 6: Individual Assignment (Variant 2) — class-level constraint rejecting reprintYear < year via #fields.globalErrors()")
add_screenshot_figure("screenshots/07_success_redirect_prg.png", "Figure 7: Successful submission, Post/Redirect/Get redirect to /books with one-time flash message banner")
add_screenshot_figure("screenshots/08_page_not_found_advice.png", "Figure 8: Browser request to GET /books/999 rendered as custom HTML error page by PageExceptionHandler")
add_screenshot_figure("screenshots/09_api_not_found_advice.png", "Figure 9: REST API request to GET /api/books/999 returning RFC 7807 ProblemDetail JSON by ApiExceptionHandler")

# ==================== 5. SECTION 4.4 TABLE ====================
add_header_paragraph("5. Task 4.4 Verification Table: One Exception, Two Audiences", level=1)
add_body_paragraph("The table below documents the behavior of BookNotFoundException across page and API controllers:")

table_44 = doc.add_table(rows=3, cols=3)
table_44.alignment = WD_TABLE_ALIGNMENT.CENTER
headers = ["Request", "Status and Content-Type", "Which Handler Ran"]
for j, h in enumerate(headers):
    c = table_44.rows[0].cells[j]
    set_cell_background(c, "1E3A8A")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(h)
    run.font.name = 'Calibri'
    run.font.size = Pt(9.5)
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

rows_data = [
    ("GET /books/999 (Browser)", "200 OK\nContent-Type: text/html;charset=UTF-8", "PageExceptionHandler.handleNotFound\n(returns template 'error/not-found')"),
    ("GET /api/books/999 (curl / HTTP Client)", "404 Not Found\nContent-Type: application/problem+json", "ApiExceptionHandler.handleNotFound\n(returns ProblemDetail JSON)")
]

for i, (req, status_ct, handler) in enumerate(rows_data):
    r = table_44.rows[i + 1]
    for j, val in enumerate([req, status_ct, handler]):
        c = r.cells[j]
        set_cell_background(c, "F8FAFC" if i % 2 == 1 else "FFFFFF")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(val)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.0)
        run.font.color.rgb = COLOR_TEXT_MAIN

# ==================== 6. VARIANT 2 EXPLANATION ====================
add_header_paragraph("6. Individual Assignment (Variant 2) Detailed Analysis", level=1)
add_body_paragraph(
    "Variant 2 specifies: «A class-level constraint: 'A reprint year must not precede the first edition year' — a global error shown by #fields.globalErrors()».\n\n"
    "Architectural Implementation Details:\n"
    "• Constraint Target: The annotation @Target(ElementType.TYPE) is placed on the class BookForm itself rather than a field, "
    "because validating reprintYear against year requires access to both properties concurrently.\n"
    "• Validation Logic: In ValidReprintYearValidator, when form.getReprintYear() is non-null, the validator checks "
    "form.getReprintYear() >= form.getYear(). If false, a ConstraintViolation is generated.\n"
    "• Error Classification: Because the violation belongs to the object rather than a specific property path, Spring MVC classifies it as a GlobalError (ObjectError), "
    "accessible via #fields.hasGlobalErrors() and iterated using #fields.errors('global').\n"
    "• Template Presentation: In form.html, an orange-accented alert box specifically surfaces global errors directly above the input fields, "
    "complementing the general summary list."
)

# ==================== 7. CONCLUSIONS ====================
add_header_paragraph("7. Conclusions", level=1)
add_body_paragraph(
    "Declarative validation using Jakarta Bean Validation and Spring MVC provides a robust, maintainable architecture compared to imperative checks in controllers. "
    "By annotating domain and form properties with standardized constraints, validation rules remain clear, reusable, and co-located with data models. "
    "Integrating BindingResult guarantees that invalid form submissions never trigger disruptive exceptions, enabling non-destructive form re-rendering with preserved user input. "
    "Furthermore, decoupling views through Spring's MessageSource enables seamless internationalization, while targeted @ControllerAdvice hierarchies establish clean separation between server-rendered web pages and REST API contracts."
)

# Save Document
output_docx = 'Report_Lab6_Forms_Validation_Variant2.docx'
doc.save(output_docx)
print(f"Report saved to {output_docx} successfully!")
